"""Ingestion pipeline — extract text, chunk, store locally, update source status.

All entry points are async and match the calls in routes/sources.py. Text
extraction for binary formats runs in a worker thread. Extraction deps
(pypdf/openpyxl) are imported lazily so a missing optional dep degrades a single
source to 'failed' rather than crashing the service.
"""

from __future__ import annotations

import asyncio
import io
import logging

from services import scrape_adapter, source_registry
from services.config import RAG_CHUNK_OVERLAP, RAG_CHUNK_SIZE

logger = logging.getLogger("muslimbot.kb.ingest")


def chunk_text(text: str, size: int = RAG_CHUNK_SIZE, overlap: int = RAG_CHUNK_OVERLAP) -> list[str]:
    text = " ".join((text or "").split())
    if not text:
        return []
    if size <= 0:
        return [text]
    step = max(1, size - overlap)
    return [text[i : i + size] for i in range(0, len(text), step)]


async def _index_text(tenant_id: str, source_id: str, title: str, text: str) -> None:
    text = (text or "").strip()
    if not text:
        await source_registry.update_source(source_id, status="failed", error="no text extracted")
        return
    chunks = chunk_text(text)
    n = await source_registry.store_chunks(tenant_id, source_id, title, chunks)
    await source_registry.update_source(
        source_id, status="indexed", chunk_count=n, error=None, rag_file_id=f"local:{source_id}"
    )
    logger.info("indexed source %s (%d chunks)", source_id, n)


async def _fail(source_id: str, message: str) -> None:
    await source_registry.update_source(source_id, status="failed", error=message[:500])


# ── text / bulk ───────────────────────────────────────────────────────────
async def ingest_text_content(
    tenant_id: str, source_id: str, title: str, content: str, filename: str = "bulk.md"
) -> None:
    await _index_text(tenant_id, source_id, title, content)


async def ingest_bulk_pages(tenant_id: str, source_id: str, title: str, pages: list) -> None:
    parts = []
    for p in pages or []:
        if isinstance(p, dict):
            parts.append(str(p.get("content") or p.get("text") or ""))
        else:
            parts.append(str(p))
    await _index_text(tenant_id, source_id, title, "\n\n".join(parts))


# ── files ─────────────────────────────────────────────────────────────────
def _extract_file(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader  # lazy
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook  # lazy
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(out)
        # csv / txt / md / json / html → best-effort decode
        return data.decode("utf-8", errors="ignore")
    except Exception as exc:  # pragma: no cover - depends on optional deps
        logger.warning("extract failed for %s: %s", filename, exc)
        return ""


async def ingest_file_bytes(
    tenant_id: str, source_id: str, filename: str, data: bytes, title: str
) -> None:
    try:
        text = await asyncio.to_thread(_extract_file, filename, data)
        await _index_text(tenant_id, source_id, title, text)
    except Exception as exc:
        await _fail(source_id, str(exc))


# ── urls / links / scrape ──────────────────────────────────────────────────
async def _ingest_web(tenant_id: str, source_id: str, title: str, url: str) -> None:
    try:
        if "youtube.com" in url or "youtu.be" in url:
            text = await asyncio.to_thread(scrape_adapter.fetch_youtube_transcript, url)
        else:
            text = await asyncio.to_thread(scrape_adapter.fetch_text, url)
        await _index_text(tenant_id, source_id, title, text)
    except Exception as exc:
        await _fail(source_id, str(exc))


async def ingest_url(
    tenant_id: str, source_id: str, title: str, url: str, depth: int
) -> None:
    await _ingest_web(tenant_id, source_id, title, url)


async def ingest_link(tenant_id: str, source_id: str, title: str, url: str) -> None:
    await _ingest_web(tenant_id, source_id, title, url)


async def ingest_scrape(
    tenant_id: str, source_id: str, title: str, url: str, depth: int
) -> None:
    # Depth crawling is a future enhancement; index the entry page for now.
    await _ingest_web(tenant_id, source_id, title, url)


async def resync_source(tenant_id: str, source_id: str) -> None:
    source = await source_registry.get_source(source_id)
    if not source:
        return
    url = source.get("url")
    title = source.get("title") or "source"
    if url:
        await _ingest_web(tenant_id, source_id, title, url)
    else:
        await source_registry.update_source(source_id, status="indexed")
